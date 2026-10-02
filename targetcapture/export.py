"""Download results as Excel (.xlsx, pictures in the sheet) or PDF (one block per result)."""
from __future__ import annotations

import os

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
_FONTS = [FONT, os.path.join(os.path.dirname(__file__), "fonts", "DejaVuSans.ttf"),
          "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/tahoma.ttf"]


def _fields(rows):
    keys = []
    for r in rows:
        for k in r:
            if not k.endswith("_image") and k not in keys:
                keys.append(k)
    return keys


def _text(v):
    if isinstance(v, (list, tuple)):
        return "; ".join(str(x) for x in v)
    return "" if v is None else v


def to_excel(rows: list[dict], image_dir: str, path: str, title: str = "Results") -> str:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from PIL import Image

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    images = [k for k in ("plate_image", "face_image", "vehicle_image", "frame_image") if any(r.get(k) for r in rows)]
    keys = _fields(rows)
    ws.append([k.replace("_", " ") for k in images + keys])
    for c in ws[1]:
        c.font = Font(bold=True)
    thumbs = os.path.join(os.path.dirname(os.path.abspath(path)), "_xl_thumbs")
    for i, r in enumerate(rows, start=2):
        ws.append([""] * len(images) + [_text(r.get(k)) for k in keys])
        ws.row_dimensions[i].height = 75
        for j, k in enumerate(images, start=1):
            if not r.get(k):
                continue
            src = os.path.join(image_dir, r[k])
            if not os.path.exists(src):
                continue
            os.makedirs(thumbs, exist_ok=True)
            im = Image.open(src)
            im.thumbnail((180, 95))
            small = os.path.join(thumbs, f"{i}_{k}.png")
            im.save(small)
            ws.add_image(XLImage(small), f"{get_column_letter(j)}{i}")
    for j in range(1, len(images) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 27
    for j, k in enumerate(keys, start=len(images) + 1):
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(40, len(k) + 4))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def to_pdf(rows: list[dict], image_dir: str, path: str, title: str = "Results", summary: str = "") -> str:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    font = next((f for f in _FONTS if os.path.exists(f)), None)
    if font:
        pdf.add_font("body", "", font)
        pdf.set_font("body", size=10)
        try:
            pdf.set_text_shaping(True)  # Arabic / Kurdish letters joined and right-to-left
        except Exception:
            pass
    else:
        pdf.set_font("helvetica", size=10)
    pdf.set_auto_page_break(True, 12)
    pdf.add_page()
    pdf.set_font_size(16)
    pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font_size(10)
    if summary:
        pdf.multi_cell(0, 6, summary, new_x="LMARGIN", new_y="NEXT")
    keys = _fields(rows)
    for r in rows:
        if pdf.get_y() > 200:
            pdf.add_page()
        pdf.ln(3)
        y = pdf.get_y()
        x = pdf.l_margin
        tallest = 0
        for k in ("frame_image", "plate_image", "face_image", "vehicle_image", "second_frame_image"):
            src = os.path.join(image_dir, r[k]) if r.get(k) else None
            if not src or not os.path.exists(src):
                continue
            w = 85 if "frame" in k else 40
            if x + w > pdf.w - pdf.r_margin:
                break
            info = pdf.image(src, x=x, y=y, w=w)
            tallest = max(tallest, info.rendered_height)
            x += w + 3
        pdf.set_y(y + tallest + 2)
        for k in keys:
            v = _text(r.get(k))
            if v == "":
                continue
            pdf.set_x(pdf.l_margin)
            pdf.cell(45, 5, k.replace("_", " "))
            pdf.multi_cell(0, 5, str(v), new_x="LMARGIN", new_y="NEXT")
        pdf.line(pdf.l_margin, pdf.get_y() + 1, pdf.w - pdf.r_margin, pdf.get_y() + 1)
    pdf.output(path)
    return path
